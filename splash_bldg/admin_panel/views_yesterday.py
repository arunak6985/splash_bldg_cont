from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth import authenticate, login, logout
from django.http import JsonResponse
from .models import JobTitle, Location, EmploymentType, AttendanceRecord, Supervisor, EmployeeAttendance
from site_application.models import JobVacancy, JobApplication
from functools import wraps

def custom_staff_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('admin_login')
        if not request.user.is_staff:
            return redirect('admin_login')
        return view_func(request, *args, **kwargs)
    return _wrapped_view

def supervisor_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.session.get('supervisor_id'):
            return redirect('admin_login')
        return view_func(request, *args, **kwargs)
    return _wrapped_view

@custom_staff_required
def admin_vacancy_management(request):
    """Admin-only view for managing job vacancies"""
    jobs = JobVacancy.objects.all().order_by('-created_at')
    
    # Calculate statistics
    active_jobs_count = jobs.filter(is_active=True).count()
    total_positions = sum(job.total_positions for job in jobs)
    available_positions = sum(job.available_positions for job in jobs)
    
    context = {
        'jobs': jobs,
        'active_jobs_count': active_jobs_count,
        'total_positions': total_positions,
        'available_positions': available_positions,
    }
    return render(request, 'admin_vacancy_management.html', context)

@custom_staff_required
def toggle_job_status(request, job_id):
    """Admin-only view to toggle job active status"""
    if request.method == 'POST':
        job = get_object_or_404(JobVacancy, id=job_id)
        job.is_active = not job.is_active
        job.save()
        status = "opened" if job.is_active else "closed"
        return JsonResponse({
            'success': True, 
            'message': f'Job "{job.title}" has been {status}.',
            'is_active': job.is_active
        })
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def update_vacancy_count(request, job_id):
    """Admin-only view to update vacancy counts"""
    if request.method == 'POST':
        job = get_object_or_404(JobVacancy, id=job_id)
        total_positions = int(request.POST.get('total_positions', job.total_positions))
        filled_positions = int(request.POST.get('filled_positions', job.filled_positions))
        
        if filled_positions > total_positions:
            return JsonResponse({
                'success': False, 
                'message': 'Filled positions cannot exceed total positions.'
            })
        
        job.total_positions = total_positions
        job.filled_positions = filled_positions
        job.save()
        
        return JsonResponse({
            'success': True,
            'message': f'Vacancy count updated for "{job.title}".',
            'available_positions': job.available_positions
        })
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def create_job_vacancy(request):
    """Admin-only view to create new job vacancy"""
    if request.method == 'POST':
        job = JobVacancy(
            title=request.POST['title'],
            description=request.POST['description'],
            requirements=request.POST['requirements'],
            location=request.POST['location'],
            salary_range=request.POST.get('salary_range', ''),
            employment_type=request.POST['employment_type'],
            total_positions=int(request.POST.get('total_positions', 1)),
            is_active=request.POST.get('is_active') == 'on'
        )
        job.save()
        messages.success(request, f'Job vacancy "{job.title}" created successfully!')
        return redirect('admin_vacancy_management')
    
    job_titles = JobTitle.objects.all()
    locations = Location.objects.all()
    jobs = JobVacancy.objects.all().order_by('-created_at')
    return render(request, 'create_job_vacancy.html', {
        'job_titles': job_titles,
        'locations': locations,
        'jobs': jobs
    })

@custom_staff_required
def edit_job_vacancy(request, job_id):
    """Admin-only view to edit job vacancy"""
    job = get_object_or_404(JobVacancy, id=job_id)
    
    if request.method == 'POST':
        job.title = request.POST['title']
        job.description = request.POST['description']
        job.requirements = request.POST['requirements']
        job.location = request.POST['location']
        job.salary_range = request.POST.get('salary_range', '')
        job.employment_type = request.POST['employment_type']
        job.total_positions = int(request.POST.get('total_positions', 1))
        job.is_active = request.POST.get('is_active') == 'on'
        job.save()
        
        return JsonResponse({
            'success': True,
            'message': f'Job vacancy "{job.title}" updated successfully!'
        })
    
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def delete_job_vacancy(request, job_id):
    """Admin-only view to delete job vacancy"""
    if request.method == 'POST':
        job = get_object_or_404(JobVacancy, id=job_id)
        job_title = job.title
        job.delete()
        
        return JsonResponse({
            'success': True,
            'message': f'Job vacancy "{job_title}" deleted successfully!'
        })
    
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def get_job_vacancy(request, job_id):
    """Admin-only view to get job vacancy details for editing"""
    job = get_object_or_404(JobVacancy, id=job_id)
    
    return JsonResponse({
        'success': True,
        'job': {
            'id': job.id,
            'title': job.title,
            'description': job.description,
            'requirements': job.requirements,
            'location': job.location,
            'salary_range': job.salary_range,
            'employment_type': job.employment_type,
            'total_positions': job.total_positions,
            'is_active': job.is_active
        }
    })

def admin_login(request):
    """Unified login page for admin and supervisors"""
    if request.user.is_authenticated and request.user.is_staff:
        return redirect('admin_vacancy_management')
    elif request.session.get('supervisor_id'):
        return redirect('supervisor_dashboard')
    
    if request.method == 'POST':
        username = request.POST['username']
        password = request.POST['password']
        
        # Check admin login first
        user = authenticate(request, username=username, password=password)
        if user is not None and user.is_staff:
            login(request, user)
            return redirect('admin_vacancy_management')
        
        # Check supervisor login
        try:
            supervisor = Supervisor.objects.get(username=username.lower(), password=password)
            request.session['supervisor_id'] = supervisor.id
            request.session['supervisor_name'] = supervisor.name
            return redirect('supervisor_dashboard')
        except Supervisor.DoesNotExist:
            pass
        
        messages.error(request, 'Invalid credentials.')
    
    return render(request, 'admin_login.html')

def admin_logout(request):
    """Admin and supervisor logout"""
    logout(request)
    if 'supervisor_id' in request.session:
        del request.session['supervisor_id']
        del request.session['supervisor_name']
    messages.success(request, 'You have been logged out successfully.')
    return redirect('admin_login')

@custom_staff_required
def job_applications(request):
    """Admin view for job applications"""
    from datetime import datetime, timedelta
    
    applications = JobApplication.objects.all().order_by('-applied_at')
    
    # Calculate statistics
    today = datetime.now().date()
    week_ago = today - timedelta(days=7)
    
    today_applications = applications.filter(applied_at__date=today).count()
    week_applications = applications.filter(applied_at__date__gte=week_ago).count()
    unique_jobs = applications.values('job_vacancy').distinct().count()
    
    context = {
        'applications': applications,
        'today_applications': today_applications,
        'week_applications': week_applications,
        'unique_jobs': unique_jobs,
    }
    return render(request, 'job_applications.html', context)

@custom_staff_required
def job_titles(request):
    job_titles = JobTitle.objects.all()
    return render(request, 'job_titles.html', {'job_titles': job_titles})

@custom_staff_required
def add_job_title(request):
    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        if title:
            job_title, created = JobTitle.objects.get_or_create(name=title)
            if created:
                return JsonResponse({'success': True, 'message': 'Job title added successfully'})
            else:
                return JsonResponse({'success': False, 'message': 'Job title already exists'})
        return JsonResponse({'success': False, 'message': 'Title is required'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def delete_job_title(request, title_id):
    if request.method == 'POST':
        try:
            job_title = get_object_or_404(JobTitle, id=title_id)
            job_title.delete()
            return JsonResponse({'success': True, 'message': 'Job title deleted successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': 'Error deleting job title'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def locations(request):
    locations = Location.objects.all()
    return render(request, 'locations.html', {'locations': locations})

@custom_staff_required
def add_location(request):
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        if name:
            location, created = Location.objects.get_or_create(name=name)
            if created:
                return JsonResponse({'success': True, 'message': 'Location added successfully'})
            else:
                return JsonResponse({'success': False, 'message': 'Location already exists'})
        return JsonResponse({'success': False, 'message': 'Location name is required'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def delete_location(request, location_id):
    if request.method == 'POST':
        try:
            location = get_object_or_404(Location, id=location_id)
            location.delete()
            return JsonResponse({'success': True, 'message': 'Location deleted successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': 'Error deleting location'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def attendance(request):
    from django.http import HttpResponse
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
    from datetime import datetime
    import calendar
    import json
    
    # Handle file upload
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            excel_file = request.FILES['excel_file']
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
            
            employees = []
            month = request.POST.get('month', 'January')
            year = int(request.POST.get('year', datetime.now().year))
            
            def parse_date(date_value):
                """Parse date from Excel cell value"""
                if not date_value:
                    return None
                
                if isinstance(date_value, datetime):
                    return date_value.date()
                elif hasattr(date_value, 'date') and callable(date_value.date):
                    return date_value.date()
                elif isinstance(date_value, str):
                    try:
                        # Clean the date string and try different formats
                        date_str = str(date_value).strip()
                        # Handle potential typos like 12025 -> 2025
                        if '/12025' in date_str:
                            date_str = date_str.replace('/12025', '/2025')
                        
                        for fmt in ['%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y']:
                            try:
                                return datetime.strptime(date_str, fmt).date()
                            except ValueError:
                                continue
                    except:
                        pass
                return None
            
            # Read data from Excel (format: REF.NO, NAME, CAT, NEW JOINING, DUTY STOP, RE JOINING, then daily attendance)
            for row in range(2, ws.max_row + 1):  # Start from row 2 (skip header)
                ref_no = ws.cell(row=row, column=1).value
                name = ws.cell(row=row, column=2).value
                cat = ws.cell(row=row, column=3).value
                new_joining = ws.cell(row=row, column=4).value
                duty_stop = ws.cell(row=row, column=5).value
                re_joining = ws.cell(row=row, column=6).value
                
                if ref_no and name:
                    # Parse date fields
                    new_joining_date = parse_date(new_joining)
                    duty_stop_date = parse_date(duty_stop)
                    re_joining_date = parse_date(re_joining)
                    
                    # Read attendance data for 31 days (columns 7-37)
                    attendance_data = {}
                    for day in range(1, 32):  # Days 1-31
                        col = day + 6  # Column 7 onwards (after date fields)
                        if col <= ws.max_column:
                            cell_value = ws.cell(row=row, column=col).value
                            attendance_data[str(day)] = str(cell_value).upper() if cell_value else ''
                    
                    # Save to database
                    AttendanceRecord.objects.update_or_create(
                        ref_no=str(ref_no).strip(),
                        month=month,
                        year=year,
                        defaults={
                            'name': str(name).strip(),
                            'category': str(cat).strip() if cat else '',
                            'new_joining': new_joining_date,
                            'duty_stop': duty_stop_date,
                            're_joining': re_joining_date,
                            'attendance_data': attendance_data
                        }
                    )
                    
                    employees.append({
                        'ref_no': str(ref_no).strip(),
                        'name': str(name).strip(),
                        'cat': str(cat).strip() if cat else '',
                        'new_joining': new_joining_date.strftime('%d/%m/%Y') if new_joining_date else '',
                        'duty_stop': duty_stop_date.strftime('%d/%m/%Y') if duty_stop_date else '',
                        're_joining': re_joining_date.strftime('%d/%m/%Y') if re_joining_date else '',
                        'attendance': attendance_data
                    })
            
            return JsonResponse({
                'success': True,
                'employees': employees,
                'count': len(employees),
                'message': f'Successfully uploaded {len(employees)} attendance records for {month} {year}'
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    
    # Handle PDF download
    if request.GET.get('download') == 'pdf':
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import Table, TableStyle
        from reportlab.lib.units import inch
        from io import BytesIO
        
        year = int(request.GET.get('year', datetime.now().year))
        month = int(request.GET.get('month', datetime.now().month))
        name = request.GET.get('name', '')
        ref_no = request.GET.get('ref_no', '')
        cat = request.GET.get('cat', '')
        
        buffer = BytesIO()
        p = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4
        
        month_names = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
                      "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
        
        # Company header
        p.setFont("Helvetica-Bold", 16)
        p.drawCentredString(width/2, height-40, "SPLASH BUILDING CONTRACTING L.L.C. - U.A.E.")
        
        # Month header
        p.setFont("Helvetica-Bold", 12)
        p.drawCentredString(width/2, height-70, f"LABOUR ATTENDANCE CARD FOR THE MONTH OF {month_names[month]} {year}")
        
        # Employee info boxes matching reference design
        p.setFont("Helvetica-Bold", 10)
        
        # Draw boxes for employee info
        p.rect(50, height-130, 250, 20)
        p.rect(300, height-130, 120, 20)
        p.rect(420, height-130, 80, 20)
        
        # Add labels and data
        p.drawString(55, height-125, f"NAME: {name}")
        p.drawString(305, height-125, f"REF.NO - {ref_no}")
        p.drawString(425, height-125, f"CAT: {cat}")
        
        # Main attendance table with integrated bottom section
        table_data = [['Date', 'P', 'OT', 'Bonus OT', 'Site No.', 'Remarks if any with Sign']]
        
        # Get attendance data from request
        attendance_data = {}
        try:
            import urllib.parse
            attendance_param = request.GET.get('attendance', '{}')
            attendance_data = json.loads(urllib.parse.unquote(attendance_param))
        except:
            pass
        
        # Days of month with Sunday marking and attendance data
        days_in_month = calendar.monthrange(year, month + 1)[1]
        for day in range(1, days_in_month + 1):
            date_obj = datetime(year, month + 1, day)
            is_sunday = date_obj.weekday() == 6
            day_str = f"{day}" if not is_sunday else f"{day} (SUN)"
            
            # Get attendance for this day
            p_val = attendance_data.get(str(day), '')
            
            table_data.append([day_str, p_val, '', '', '', ''])
        
        table_data.append(['Total', '', '', '', '', ''])
        
        # Add integrated bottom section rows
        table_data.append(["Engineer's Sign", "Employee Sign", "No of Days :", "", "", "Basic"])
        table_data.append(["", "", "Normal OT :", "", "", "OT"])
        table_data.append(["", "", "Bonus OT :", "", "", "Bonus"])
        table_data.append(["", "", "", "", "", "Gross Salary"])
        table_data.append(["", "", "", "", "", "Adv Deduction"])
        table_data.append(["", "", "", "", "", "Net Salary"])
        
        # Create table with adjusted column widths
        table = Table(table_data, colWidths=[0.8*inch, 0.6*inch, 0.8*inch, 1*inch, 1*inch, 2.5*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.beige),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
            ('BACKGROUND', (0, 1), (-1, days_in_month), colors.lightyellow),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, days_in_month+1), (-1, days_in_month+1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            # Span cells for signature boxes
            ('SPAN', (0, days_in_month+2), (0, days_in_month+7)),
            ('SPAN', (1, days_in_month+2), (1, days_in_month+7)),
        ]))
        
        # Mark Sundays and Absent days in red
        for i in range(1, days_in_month + 1):
            row_index = i
            day_str = table_data[row_index][0]
            attendance_val = attendance_data.get(str(i), '')
            
            if '(SUN)' in str(day_str):
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, row_index), (0, row_index), colors.red),
                    ('TEXTCOLOR', (0, row_index), (0, row_index), colors.white),
                ]))
            # Mark entire row red if attendance is 'A' (Absent)
            if str(attendance_val).upper() == 'A':
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, row_index), (-1, row_index), colors.red),
                    ('TEXTCOLOR', (0, row_index), (-1, row_index), colors.white),
                ]))
        
        # Draw table
        table.wrapOn(p, width, height)
        table.drawOn(p, 50, height-550)
        
        p.showPage()
        p.save()
        
        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="attendance_{ref_no}_{month_names[month]}_{year}.pdf"'
        return response
    
    # Handle bulk PDF generation
    if request.method == 'POST' and request.POST.get('generate_pdfs'):
        import zipfile
        from io import BytesIO
        import json
        
        try:
            employees_data = json.loads(request.POST.get('employees_data', '[]'))
            year = int(request.POST.get('year', datetime.now().year))
            month = int(request.POST.get('month', datetime.now().month))
            
            # Create ZIP file
            zip_buffer = BytesIO()
            with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                
                for emp in employees_data:
                    # Generate PDF for each employee
                    pdf_buffer = BytesIO()
                    p = canvas.Canvas(pdf_buffer, pagesize=A4)
                    width, height = A4
                    
                    month_names = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
                                  "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
                    
                    # Company header
                    p.setFont("Helvetica-Bold", 16)
                    p.drawCentredString(width/2, height-40, "SPLASH BUILDING CONTRACTING L.L.C. - U.A.E.")
                    
                    # Month header
                    p.setFont("Helvetica-Bold", 12)
                    p.drawCentredString(width/2, height-70, f"LABOUR ATTENDANCE CARD FOR THE MONTH OF {month_names[month]} {year}")
                    
                    # Employee info boxes matching reference design
                    p.setFont("Helvetica-Bold", 10)
                    
                    # Draw boxes for employee info
                    p.rect(50, height-130, 250, 20)
                    p.rect(300, height-130, 120, 20)
                    p.rect(420, height-130, 80, 20)
                    
                    # Add labels and data
                    p.drawString(55, height-125, f"NAME: {emp.get('name', '')}")
                    p.drawString(305, height-125, f"REF.NO - {emp.get('ref_no', '')}")
                    p.drawString(425, height-125, f"CAT: {emp.get('cat', '')}")
                    
                    # Table data
                    table_data = [['Date', 'P', 'OT', 'Bonus OT', 'Site No.', 'Remarks if any with Sign']]
                    
                    # Days of month with Sunday marking and attendance data
                    days_in_month = calendar.monthrange(year, month + 1)[1]
                    holiday_rows = []
                    medical_rows = []
                    absent_rows = []
                    
                    for day in range(1, days_in_month + 1):
                        date_obj = datetime(year, month + 1, day)
                        is_sunday = date_obj.weekday() == 6
                        day_str = f"{day}" if not is_sunday else f"{day} (SUN)"
                        
                        # Get attendance for this day
                        attendance_val = emp.get('attendance', {}).get(str(day), '')
                        
                        # Handle special attendance codes
                        if str(attendance_val).upper() == 'H':
                            # Holiday - merge columns with HOLIDAY text
                            table_data.append([day_str, 'HOLIDAY', '', '', '', ''])
                            holiday_rows.append(day)
                        elif str(attendance_val).upper() == 'M':
                            # Medical - merge columns with MEDICAL LEAVE text
                            table_data.append([day_str, 'MEDICAL LEAVE', '', '', '', ''])
                            medical_rows.append(day)
                        elif str(attendance_val).upper() == 'A':
                            # Absent - red dashes
                            table_data.append([day_str, '-', '-', '-', '-', ''])
                            absent_rows.append(day)
                        else:
                            # Regular attendance
                            table_data.append([day_str, attendance_val, '', '', '', ''])
                    
                    table_data.append(['Total', '', '', '', '', ''])
                    
                    # Create table
                    table = Table(table_data, colWidths=[0.8*inch, 0.6*inch, 0.8*inch, 1*inch, 1*inch, 2.5*inch])
                    table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.beige),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                        ('BACKGROUND', (0, 1), (-1, days_in_month), colors.lightyellow),
                        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, -1), 8),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                        ('BACKGROUND', (0, -1), (-1, -1), colors.beige),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black),
                    ]))
                    
                    # Mark Sundays, Absent, Holiday, and Medical days
                    for i in range(1, days_in_month + 1):
                        row_index = i
                        day_str = table_data[row_index][0]
                        
                        if '(SUN)' in str(day_str):
                            table.setStyle(TableStyle([
                                ('BACKGROUND', (0, row_index), (0, row_index), colors.red),
                                ('TEXTCOLOR', (0, row_index), (0, row_index), colors.white),
                            ]))
                        
                        # Mark absent days in red
                        if i in absent_rows:
                            table.setStyle(TableStyle([
                                ('TEXTCOLOR', (1, row_index), (4, row_index), colors.red),
                                ('FONTNAME', (1, row_index), (4, row_index), 'Helvetica-Bold'),
                            ]))
                        
                        # Merge and style Holiday rows - yellow text, no background
                        if i in holiday_rows:
                            table.setStyle(TableStyle([
                                ('SPAN', (1, row_index), (4, row_index)),
                                ('TEXTCOLOR', (1, row_index), (4, row_index), colors.yellow),
                                ('FONTNAME', (1, row_index), (4, row_index), 'Helvetica-Bold'),
                                ('ALIGN', (1, row_index), (4, row_index), 'CENTER'),
                            ]))
                        
                        # Merge and style Medical rows - red text, no background
                        if i in medical_rows:
                            table.setStyle(TableStyle([
                                ('SPAN', (1, row_index), (4, row_index)),
                                ('TEXTCOLOR', (1, row_index), (4, row_index), colors.red),
                                ('FONTNAME', (1, row_index), (4, row_index), 'Helvetica-Bold'),
                                ('ALIGN', (1, row_index), (4, row_index), 'CENTER'),
                            ]))
                    
                    table.wrapOn(p, width, height)
                    table.drawOn(p, 50, height-520)
                    
                    # Bottom section matching reference design (small gap after Total)
                    bottom_y = 160
                    
                    # Engineer's Sign box
                    p.rect(50, bottom_y-100, 120, 80)
                    p.setFont("Helvetica-Bold", 9)
                    p.drawString(55, bottom_y-15, "Engineer's Sign")
                    
                    # Employee Sign box
                    p.rect(180, bottom_y-100, 120, 80)
                    p.drawString(185, bottom_y-15, "Employee Sign")
                    
                    # Middle section with calculations
                    p.drawString(320, bottom_y-25, "No of Days :")
                    p.drawString(320, bottom_y-45, "Normal OT :")
                    p.drawString(320, bottom_y-65, "Bonus OT :")
                    
                    # Right section with salary details in boxes
                    salary_x = 450
                    box_height = 15
                    
                    # Draw salary boxes
                    p.rect(salary_x, bottom_y-20, 80, box_height)
                    p.drawString(salary_x+5, bottom_y-15, "Basic")
                    
                    p.rect(salary_x, bottom_y-35, 80, box_height)
                    p.drawString(salary_x+5, bottom_y-30, "OT")
                    
                    p.rect(salary_x, bottom_y-50, 80, box_height)
                    p.drawString(salary_x+5, bottom_y-45, "Bonus")
                    
                    p.rect(salary_x, bottom_y-65, 80, box_height)
                    p.drawString(salary_x+5, bottom_y-60, "Gross Salary")
                    
                    p.rect(salary_x, bottom_y-80, 80, box_height)
                    p.drawString(salary_x+5, bottom_y-75, "Adv Deduction")
                    
                    p.rect(salary_x, bottom_y-95, 80, box_height)
                    p.drawString(salary_x+5, bottom_y-90, "Net Salary")
                    
                    p.showPage()
                    p.save()
                    
                    # Add PDF to ZIP
                    pdf_buffer.seek(0)
                    filename = f"attendance_{emp.get('ref_no', 'unknown')}_{month_names[month]}_{year}.pdf"
                    zip_file.writestr(filename, pdf_buffer.getvalue())
            
            zip_buffer.seek(0)
            response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
            response['Content-Disposition'] = f'attachment; filename="attendance_cards_{month_names[month]}_{year}.zip"'
            return response
            
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    
    # Handle Excel download
    if request.GET.get('download') == 'excel':
        # Create Excel file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Card"
        
        # Company header
        ws.merge_cells('A1:G1')
        ws['A1'] = 'SPLASH BUILDING CONTRACTING L.L.C. - U.A.E.'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Month header
        current_month = datetime.now().strftime('%B %Y')
        ws.merge_cells('A2:G2')
        ws['A2'] = f'LABOUR ATTENDANCE CARD FOR THE MONTH OF {current_month.upper()}'
        ws['A2'].font = Font(bold=True, size=12)
        ws['A2'].alignment = Alignment(horizontal='center')
        
        # Employee info row
        ws.merge_cells('A3:D3')
        ws['A3'] = 'NAME:'
        ws.merge_cells('E3:F3')
        ws['E3'] = 'REF.NO -'
        ws['G3'] = 'CAT:'
        
        # Table headers
        headers = ['Date', 'P', 'OT', 'Bonus OT', 'Site No.', 'Remarks if any with Sign']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        
        # Days of month
        current_date = datetime.now()
        days_in_month = calendar.monthrange(current_date.year, current_date.month)[1]
        
        for day in range(1, days_in_month + 1):
            row = day + 4
            ws.cell(row=row, column=1, value=day)
            for col in range(2, 7):
                ws.cell(row=row, column=col, value='')
        
        # Total row
        total_row = days_in_month + 5
        ws.cell(row=total_row, column=1, value='Total')
        
        # Bottom section
        bottom_start = total_row + 2
        ws.merge_cells(f'A{bottom_start}:B{bottom_start+2}')
        ws[f'A{bottom_start}'] = "Engineer's Sign"
        
        ws.merge_cells(f'C{bottom_start}:D{bottom_start+2}')
        ws[f'C{bottom_start}'] = "Employee Sign"
        
        ws[f'E{bottom_start}'] = 'No of Days :'
        ws[f'E{bottom_start+1}'] = 'Normal OT :'
        ws[f'E{bottom_start+2}'] = 'Bonus OT :'
        
        ws[f'F{bottom_start}'] = 'Basic'
        ws[f'F{bottom_start+1}'] = 'OT'
        ws[f'F{bottom_start+2}'] = 'Bonus'
        ws[f'F{bottom_start+3}'] = 'Gross Salary'
        ws[f'F{bottom_start+4}'] = 'Adv Deduction'
        ws[f'F{bottom_start+5}'] = 'Net Salary'
        
        # Add borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in ws.iter_rows(min_row=3, max_row=total_row+6, min_col=1, max_col=6):
            for cell in row:
                cell.border = thin_border
        
        # Set column widths
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 5
        ws.column_dimensions['C'].width = 8
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 25
        
        year = int(request.GET.get('year', datetime.now().year))
        month = int(request.GET.get('month', datetime.now().month))
        
        month_names = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
                      "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
        selected_month = f"{month_names[month]} {year}"
        ws['A2'] = f'LABOUR ATTENDANCE CARD FOR THE MONTH OF {selected_month}'
        
        # Update days for selected month
        days_in_month = calendar.monthrange(year, month + 1)[1]
        
        # Clear existing days and add new ones
        for row in range(5, 50):
            for col in range(1, 7):
                ws.cell(row=row, column=col).value = None
        
        for day in range(1, days_in_month + 1):
            row = day + 4
            ws.cell(row=row, column=1, value=day)
            
            # Mark Sundays in red
            date_obj = datetime(year, month + 1, day)
            if date_obj.weekday() == 6:  # Sunday
                from openpyxl.styles import PatternFill
                red_fill = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')
                ws.cell(row=row, column=1).fill = red_fill
        
        # Add borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        total_row = days_in_month + 5
        ws.cell(row=total_row, column=1, value='Total')
        
        for row in ws.iter_rows(min_row=3, max_row=total_row+6, min_col=1, max_col=6):
            for cell in row:
                cell.border = thin_border
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="attendance_card_{selected_month.replace(" ", "_")}.xlsx"'
        wb.save(response)
        return response
    
    # Get attendance records for display
    selected_month = request.GET.get('month', '')
    selected_year = request.GET.get('year', '')
    
    records = AttendanceRecord.objects.all()
    if selected_month:
        records = records.filter(month=selected_month)
    if selected_year:
        records = records.filter(year=int(selected_year))
    
    context = {
        'records': records,
        'selected_month': selected_month,
        'selected_year': selected_year,
        'months': ['January', 'February', 'March', 'April', 'May', 'June',
                  'July', 'August', 'September', 'October', 'November', 'December'],
        'years': range(2020, 2030),
        'days_range': range(1, 32)
    }
    return render(request, 'attendance.html', context)

@custom_staff_required
def delete_attendance_record(request, record_id):
    if request.method == 'POST':
        try:
            record = get_object_or_404(AttendanceRecord, id=record_id)
            record.delete()
            return JsonResponse({'success': True, 'message': 'Record deleted successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def delete_attendance_bulk(request):
    if request.method == 'POST':
        try:
            import json
            data = json.loads(request.body)
            record_ids = data.get('ids', [])
            
            if not record_ids:
                return JsonResponse({'success': False, 'message': 'No records selected'})
            
            deleted_count = AttendanceRecord.objects.filter(id__in=record_ids).delete()[0]
            return JsonResponse({
                'success': True, 
                'message': f'Successfully deleted {deleted_count} records'
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def supervisors(request):
    supervisors = Supervisor.objects.all()
    return render(request, 'supervisors.html', {'supervisors': supervisors})

@custom_staff_required
def add_supervisor(request):
    if request.method == 'POST':
        try:
            employee_id = request.POST.get('employee_id', '').strip()
            username = request.POST.get('username', '').strip().lower()
            password = request.POST.get('password', '').strip()
            
            # Validate employee ID uniqueness
            if Supervisor.objects.filter(employee_id=employee_id).exists():
                return JsonResponse({'success': False, 'message': 'Employee ID already exists'})
            
            # Validate username format and uniqueness
            if not username.endswith('@gmail.com'):
                return JsonResponse({'success': False, 'message': 'Username must end with @gmail.com'})
            if Supervisor.objects.filter(username=username).exists():
                return JsonResponse({'success': False, 'message': 'Username already exists'})
            
            # Validate password format
            import re
            if not re.match(r'^[A-Z][a-z]+[@#$%^&*][0-9]+$', password):
                return JsonResponse({'success': False, 'message': 'Password must start with capital letter, followed by lowercase letters, then a symbol (@#$%^&*), then numbers'})
            
            supervisor = Supervisor(
                employee_id=employee_id,
                name=request.POST.get('name', '').strip(),
                mobile_number=request.POST.get('mobile_number', '').strip(),
                site_location=request.POST.get('site_location', '').strip(),
                username=username,
                password=password
            )
            supervisor.save()
            return JsonResponse({'success': True, 'message': 'Supervisor added successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def get_supervisor(request, supervisor_id):
    supervisor = get_object_or_404(Supervisor, id=supervisor_id)
    return JsonResponse({
        'success': True,
        'supervisor': {
            'id': supervisor.id,
            'employee_id': supervisor.employee_id,
            'name': supervisor.name,
            'mobile_number': supervisor.mobile_number,
            'site_location': supervisor.site_location,
            'username': supervisor.username,
            'password': supervisor.password
        }
    })

@custom_staff_required
def edit_supervisor(request, supervisor_id):
    if request.method == 'POST':
        try:
            supervisor = get_object_or_404(Supervisor, id=supervisor_id)
            employee_id = request.POST.get('employee_id', '').strip()
            username = request.POST.get('username', '').strip().lower()
            password = request.POST.get('password', '').strip()
            
            # Validate employee ID uniqueness (exclude current supervisor)
            if Supervisor.objects.filter(employee_id=employee_id).exclude(id=supervisor_id).exists():
                return JsonResponse({'success': False, 'message': 'Employee ID already exists'})
            
            # Validate username format and uniqueness (exclude current supervisor)
            if not username.endswith('@gmail.com'):
                return JsonResponse({'success': False, 'message': 'Username must end with @gmail.com'})
            if Supervisor.objects.filter(username=username).exclude(id=supervisor_id).exists():
                return JsonResponse({'success': False, 'message': 'Username already exists'})
            
            # Validate password format
            import re
            if not re.match(r'^[A-Z][a-z]+[@#$%^&*][0-9]+$', password):
                return JsonResponse({'success': False, 'message': 'Password must start with capital letter, followed by lowercase letters, then a symbol (@#$%^&*), then numbers'})
            
            supervisor.employee_id = employee_id
            supervisor.name = request.POST.get('name', '').strip()
            supervisor.mobile_number = request.POST.get('mobile_number', '').strip()
            supervisor.site_location = request.POST.get('site_location', '').strip()
            supervisor.username = username
            supervisor.password = password
            supervisor.save()
            
            return JsonResponse({'success': True, 'message': 'Supervisor updated successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def check_employee_id(request):
    employee_id = request.GET.get('employee_id', '').strip()
    supervisor_id = request.GET.get('supervisor_id', '')
    
    query = Supervisor.objects.filter(employee_id=employee_id)
    if supervisor_id:
        query = query.exclude(id=supervisor_id)
    
    exists = query.exists()
    return JsonResponse({'exists': exists})

@custom_staff_required
def check_username(request):
    username = request.GET.get('username', '').strip().lower()
    supervisor_id = request.GET.get('supervisor_id', '')
    
    query = Supervisor.objects.filter(username=username)
    if supervisor_id:
        query = query.exclude(id=supervisor_id)
    
    exists = query.exists()
    return JsonResponse({'exists': exists})

@custom_staff_required
def delete_supervisor(request, supervisor_id):
    if request.method == 'POST':
        try:
            supervisor = get_object_or_404(Supervisor, id=supervisor_id)
            supervisor.delete()
            return JsonResponse({'success': True, 'message': 'Supervisor deleted successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': 'Error deleting supervisor'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@supervisor_required
def supervisor_dashboard(request):
    """Supervisor dashboard with attendance only"""
    from datetime import datetime
    import calendar
    import json
    
    # Get attendance records for display
    selected_month = request.GET.get('month', '')
    selected_year = request.GET.get('year', '')
    
    records = AttendanceRecord.objects.all()
    if selected_month:
        records = records.filter(month=selected_month)
    if selected_year:
        records = records.filter(year=int(selected_year))
    
    context = {
        'records': records,
        'selected_month': selected_month,
        'selected_year': selected_year,
        'months': ['January', 'February', 'March', 'April', 'May', 'June',
                  'July', 'August', 'September', 'October', 'November', 'December'],
        'years': range(2020, 2030),
        'days_range': range(1, 32),
        'supervisor_name': request.session.get('supervisor_name', 'Supervisor')
    }
    return render(request, 'supervisor_dashboard.html', context)

@custom_staff_required
def mark_attendance(request):
    """Mark attendance page with Excel upload and monthly view"""
    import openpyxl
    from datetime import datetime, date
    import calendar
    
    def parse_date(date_value):
        """Parse date from Excel cell value"""
        if not date_value:
            return None
        
        if isinstance(date_value, datetime):
            return date_value.date()
        elif isinstance(date_value, date):
            return date_value
        elif isinstance(date_value, str):
            try:
                # Try different date formats
                for fmt in ['%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y']:
                    try:
                        return datetime.strptime(date_value.strip(), fmt).date()
                    except ValueError:
                        continue
            except:
                pass
        return None
    
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            excel_file = request.FILES['excel_file']
            month = request.POST.get('month')
            year = int(request.POST.get('year'))
            
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
            
            employees = []
            
            # Read data from Excel (assuming headers: REF NO, Name, CATEGORY, JOINING DATE, RE JOINING)
            for row in range(2, ws.max_row + 1):
                ref_no = ws.cell(row=row, column=1).value  # REF NO
                name = ws.cell(row=row, column=2).value    # Name
                category = ws.cell(row=row, column=3).value # Category
                joining_date = ws.cell(row=row, column=4).value  # Joining Date
                rejoining_date = ws.cell(row=row, column=5).value # Re-joining Date
                
                if ref_no and name:
                    # Parse dates
                    parsed_joining_date = parse_date(joining_date)
                    parsed_rejoining_date = parse_date(rejoining_date)
                    
                    # Create empty attendance data for the month
                    attendance_data = {}
                    days_in_month = calendar.monthrange(year, datetime.strptime(month, '%B').month)[1]
                    
                    for day in range(1, days_in_month + 1):
                        attendance_data[str(day)] = {'attendance': '', 'overtime': ''}
                    
                    # Save to database
                    EmployeeAttendance.objects.update_or_create(
                        ref_no=str(ref_no).strip(),
                        month=month,
                        year=year,
                        defaults={
                            'name': str(name).strip(),
                            'category': str(category).strip() if category else '',
                            'joining_date': parsed_joining_date,
                            'rejoining_date': parsed_rejoining_date,
                            'attendance_data': attendance_data
                        }
                    )
                    
                    employees.append({
                        'ref_no': str(ref_no).strip(),
                        'name': str(name).strip(),
                        'category': str(category).strip() if category else '',
                        'joining_date': parsed_joining_date.strftime('%d/%m/%Y') if parsed_joining_date else '',
                        'rejoining_date': parsed_rejoining_date.strftime('%d/%m/%Y') if parsed_rejoining_date else '',
                        'attendance': attendance_data
                    })
            
            return JsonResponse({
                'success': True,
                'employees': employees,
                'count': len(employees)
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    
    # Handle save attendance data
    if request.method == 'POST' and 'save_attendance' in request.POST:
        try:
            import json
            attendance_data = json.loads(request.POST.get('attendance_data', '{}'))
            
            for ref_no, data in attendance_data.items():
                try:
                    employee = EmployeeAttendance.objects.get(
                        ref_no=ref_no,
                        month=data['month'],
                        year=int(data['year'])
                    )
                    employee.attendance_data = data['attendance']
                    employee.save()
                except EmployeeAttendance.DoesNotExist:
                    continue
            
            return JsonResponse({'success': True, 'message': 'Attendance saved successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    
    # Get attendance records for display
    selected_month = request.GET.get('month', '')
    selected_year = request.GET.get('year', '')
    
    records = EmployeeAttendance.objects.all()
    if selected_month:
        records = records.filter(month=selected_month)
    if selected_year:
        records = records.filter(year=int(selected_year))
    
    context = {
        'records': records,
        'selected_month': selected_month,
        'selected_year': selected_year,
        'months': ['January', 'February', 'March', 'April', 'May', 'June',
                  'July', 'August', 'September', 'October', 'November', 'December'],
        'years': range(2020, 2030),
    }
    return render(request, 'mark_attendance.html', context)