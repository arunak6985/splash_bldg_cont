from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth import authenticate, login, logout
from django.http import JsonResponse, HttpResponse
from django.db import models
from .models import JobTitle, Location, EmploymentType, AttendanceRecord
from site_application.models import JobVacancy, JobApplication
from functools import wraps
import pandas as pd
import calendar
from datetime import datetime, date

def custom_staff_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('admin_login')
        if not request.user.is_staff:
            return redirect('admin_login')
        return view_func(request, *args, **kwargs)
    return _wrapped_view

@custom_staff_required
def admin_vacancy_management(request):
    """Admin-only view for managing job vacancies"""
    jobs = JobVacancy.objects.all().order_by('-created_at')
    
    # Calculate statistics
    active_jobs_count = jobs.filter(is_active=True).count()
    total_positions = sum(job.total_positions for job in jobs)
    available_positions = sum(job.available_positions for job in jobs)
    
    context = {
        'jobs': jobs,
        'active_jobs_count': active_jobs_count,
        'total_positions': total_positions,
        'available_positions': available_positions,
    }
    return render(request, 'admin_vacancy_management.html', context)

@custom_staff_required
def toggle_job_status(request, job_id):
    """Admin-only view to toggle job active status"""
    if request.method == 'POST':
        job = get_object_or_404(JobVacancy, id=job_id)
        job.is_active = not job.is_active
        job.save()
        status = "opened" if job.is_active else "closed"
        return JsonResponse({
            'success': True, 
            'message': f'Job "{job.title}" has been {status}.',
            'is_active': job.is_active
        })
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def update_vacancy_count(request, job_id):
    """Admin-only view to update vacancy counts"""
    if request.method == 'POST':
        job = get_object_or_404(JobVacancy, id=job_id)
        total_positions = int(request.POST.get('total_positions', job.total_positions))
        filled_positions = int(request.POST.get('filled_positions', job.filled_positions))
        
        if filled_positions > total_positions:
            return JsonResponse({
                'success': False, 
                'message': 'Filled positions cannot exceed total positions.'
            })
        
        job.total_positions = total_positions
        job.filled_positions = filled_positions
        job.save()
        
        return JsonResponse({
            'success': True,
            'message': f'Vacancy count updated for "{job.title}".',
            'available_positions': job.available_positions
        })
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def create_job_vacancy(request):
    """Admin-only view to create new job vacancy"""
    if request.method == 'POST':
        job = JobVacancy(
            title=request.POST['title'],
            description=request.POST['description'],
            requirements=request.POST['requirements'],
            location=request.POST['location'],
            salary_range=request.POST.get('salary_range', ''),
            employment_type=request.POST['employment_type'],
            total_positions=int(request.POST.get('total_positions', 1)),
            is_active=request.POST.get('is_active') == 'on'
        )
        job.save()
        messages.success(request, f'Job vacancy "{job.title}" created successfully!')
        return redirect('admin_vacancy_management')
    
    job_titles = JobTitle.objects.all()
    locations = Location.objects.all()
    jobs = JobVacancy.objects.all().order_by('-created_at')
    return render(request, 'create_job_vacancy.html', {
        'job_titles': job_titles,
        'locations': locations,
        'jobs': jobs
    })

@custom_staff_required
def edit_job_vacancy(request, job_id):
    """Admin-only view to edit job vacancy"""
    job = get_object_or_404(JobVacancy, id=job_id)
    
    if request.method == 'POST':
        job.title = request.POST['title']
        job.description = request.POST['description']
        job.requirements = request.POST['requirements']
        job.location = request.POST['location']
        job.salary_range = request.POST.get('salary_range', '')
        job.employment_type = request.POST['employment_type']
        job.total_positions = int(request.POST.get('total_positions', 1))
        job.is_active = request.POST.get('is_active') == 'on'
        job.save()
        
        return JsonResponse({
            'success': True,
            'message': f'Job vacancy "{job.title}" updated successfully!'
        })
    
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def delete_job_vacancy(request, job_id):
    """Admin-only view to delete job vacancy"""
    if request.method == 'POST':
        job = get_object_or_404(JobVacancy, id=job_id)
        job_title = job.title
        job.delete()
        
        return JsonResponse({
            'success': True,
            'message': f'Job vacancy "{job_title}" deleted successfully!'
        })
    
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def get_job_vacancy(request, job_id):
    """Admin-only view to get job vacancy details for editing"""
    job = get_object_or_404(JobVacancy, id=job_id)
    
    return JsonResponse({
        'success': True,
        'job': {
            'id': job.id,
            'title': job.title,
            'description': job.description,
            'requirements': job.requirements,
            'location': job.location,
            'salary_range': job.salary_range,
            'employment_type': job.employment_type,
            'total_positions': job.total_positions,
            'is_active': job.is_active
        }
    })

def admin_login(request):
    """Admin login page"""
    if request.user.is_authenticated and request.user.is_staff:
        return redirect('admin_vacancy_management')
    
    if request.method == 'POST':
        username = request.POST['username']
        password = request.POST['password']
        
        user = authenticate(request, username=username, password=password)
        if user is not None and user.is_staff:
            login(request, user)
            return redirect('admin_vacancy_management')
        
        messages.error(request, 'Invalid credentials.')
    
    return render(request, 'admin_login.html')

def admin_logout(request):
    """Admin logout"""
    logout(request)
    messages.success(request, 'You have been logged out successfully.')
    return redirect('admin_login')

@custom_staff_required
def job_applications(request):
    """Admin view for job applications"""
    from datetime import datetime, timedelta
    
    applications = JobApplication.objects.all().order_by('-applied_at')
    
    # Calculate statistics
    today = datetime.now().date()
    week_ago = today - timedelta(days=7)
    
    today_applications = applications.filter(applied_at__date=today).count()
    week_applications = applications.filter(applied_at__date__gte=week_ago).count()
    unique_jobs = applications.values('job_vacancy').distinct().count()
    
    context = {
        'applications': applications,
        'today_applications': today_applications,
        'week_applications': week_applications,
        'unique_jobs': unique_jobs,
    }
    return render(request, 'job_applications.html', context)

@custom_staff_required
def job_titles(request):
    job_titles = JobTitle.objects.all()
    return render(request, 'job_titles.html', {'job_titles': job_titles})

@custom_staff_required
def add_job_title(request):
    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        if title:
            job_title, created = JobTitle.objects.get_or_create(name=title)
            if created:
                return JsonResponse({'success': True, 'message': 'Job title added successfully'})
            else:
                return JsonResponse({'success': False, 'message': 'Job title already exists'})
        return JsonResponse({'success': False, 'message': 'Title is required'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def delete_job_title(request, title_id):
    if request.method == 'POST':
        try:
            job_title = get_object_or_404(JobTitle, id=title_id)
            job_title.delete()
            return JsonResponse({'success': True, 'message': 'Job title deleted successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': 'Error deleting job title'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def locations(request):
    locations = Location.objects.all()
    return render(request, 'locations.html', {'locations': locations})

@custom_staff_required
def add_location(request):
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        if name:
            location, created = Location.objects.get_or_create(name=name)
            if created:
                return JsonResponse({'success': True, 'message': 'Location added successfully'})
            else:
                return JsonResponse({'success': False, 'message': 'Location already exists'})
        return JsonResponse({'success': False, 'message': 'Location name is required'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def delete_location(request, location_id):
    if request.method == 'POST':
        try:
            location = get_object_or_404(Location, id=location_id)
            location.delete()
            return JsonResponse({'success': True, 'message': 'Location deleted successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': 'Error deleting location'})
    return JsonResponse({'success': False, 'message': 'Invalid request'})

@custom_staff_required
def attendance(request):
    import openpyxl
    import json
    
    # Handle file upload
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            excel_file = request.FILES['excel_file']
            month = request.POST.get('month', 'December')
            year = int(request.POST.get('year', 2025))
            clear_existing = request.POST.get('clear_existing') == '1'
            
            # Clear existing records if requested
            if clear_existing:
                AttendanceRecord.objects.filter(month=month, year=year).delete()
            
            # Read Excel file using openpyxl
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
            employees_processed = 0
            
            # Process each row starting from row 2 (skip header)
            for row_num in range(2, ws.max_row + 1):
                try:
                    # Read basic employee data from your Excel format
                    ref_no = ws.cell(row=row_num, column=1).value  # REF.NO
                    name = ws.cell(row=row_num, column=2).value    # NAME
                    category = ws.cell(row=row_num, column=3).value # CATEGORY
                    new_joining = ws.cell(row=row_num, column=4).value # NEW JOINING
                    duty_stop = ws.cell(row=row_num, column=5).value   # DUTY STOP
                    re_joining = ws.cell(row=row_num, column=6).value  # RE JOINING
                    
                    # Skip completely empty rows (both ref_no and name are empty)
                    if not ref_no and not name:
                        continue
                    
                    # Clean data - handle empty ref_no
                    ref_no = str(ref_no).strip() if ref_no else ''
                    name = str(name).strip() if name else ''
                    category = str(category).strip() if category else ''
                    
                    # Parse date fields
                    def parse_date(date_val):
                        if not date_val or str(date_val).strip() == '':
                            return None
                        try:
                            if isinstance(date_val, datetime):
                                return date_val.date()
                            date_str = str(date_val).strip()
                            # Handle date formats like 10/10/12025 -> 10/10/2025
                            if '/12025' in date_str:
                                date_str = date_str.replace('/12025', '/2025')
                            # Try different date formats
                            for fmt in ['%d/%m/%Y', '%m/%d/%Y', '%Y-%m-%d', '%d-%m-%Y']:
                                try:
                                    return datetime.strptime(date_str, fmt).date()
                                except ValueError:
                                    continue
                        except:
                            pass
                        return None
                    
                    new_joining_date = parse_date(new_joining)
                    duty_stop_date = parse_date(duty_stop)
                    re_joining_date = parse_date(re_joining)
                    
                    # Process attendance data (columns 7 onwards for days 1-31)
                    attendance_data = {}
                    days_in_month = calendar.monthrange(year, ['January', 'February', 'March', 'April', 'May', 'June',
                                                               'July', 'August', 'September', 'October', 'November', 'December'].index(month) + 1)[1]
                    
                    # Read attendance for each day (starting from column 7)
                    for day in range(1, days_in_month + 1):
                        col_num = 6 + day  # Column 7, 8, 9, etc. for days 1, 2, 3, etc.
                        if col_num <= ws.max_column:
                            cell_value = ws.cell(row=row_num, column=col_num).value
                            if cell_value is not None and str(cell_value).strip() != '':
                                attendance_data[str(day)] = str(cell_value).strip().upper()
                            else:
                                attendance_data[str(day)] = ''
                        else:
                            attendance_data[str(day)] = ''
                    
                    # Save to database - use unique identifier for empty ref_no to avoid conflicts
                    unique_ref = ref_no if ref_no else f"EMPTY_REF_{row_num}_{month}_{year}"
                    
                    AttendanceRecord.objects.update_or_create(
                        ref_no=unique_ref,
                        month=month,
                        year=year,
                        defaults={
                            'name': name or '',
                            'category': category,
                            'new_joining': new_joining_date,
                            'duty_stop': duty_stop_date,
                            're_joining': re_joining_date,
                            'attendance_data': attendance_data
                        }
                    )
                    employees_processed += 1
                    
                except Exception as e:
                    print(f"Error processing row {row_num}: {str(e)}")
                    continue
            
            return JsonResponse({
                'success': True,
                'message': f'Successfully processed {employees_processed} employees for {month} {year}',
                'count': employees_processed
            })
            
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'Error processing file: {str(e)}'})
    
    # Get attendance records for display
    month_filter = request.GET.get('month')
    year_filter = request.GET.get('year')
    
    if month_filter and year_filter:
        from datetime import date
        
        # Get month number
        month_num = ['January', 'February', 'March', 'April', 'May', 'June',
                    'July', 'August', 'September', 'October', 'November', 'December'].index(month_filter) + 1
        
        # Get first and last day of the selected month
        first_day = date(int(year_filter), month_num, 1)
        last_day = date(int(year_filter), month_num, calendar.monthrange(int(year_filter), month_num)[1])
        
        # Get all records for the month/year OR records with duty_stop in the selected month
        records = AttendanceRecord.objects.filter(
            models.Q(month=month_filter, year=int(year_filter)) |
            models.Q(duty_stop__gte=first_day, duty_stop__lte=last_day)
        ).order_by('ref_no')
    else:
        # Show all records if no filter
        records = AttendanceRecord.objects.all().order_by('-id')
    
    # Calculate attendance statistics
    total_present = 0
    total_absent = 0
    total_holiday = 0
    total_medical = 0
    total_normal_ot = 0
    total_bonus_ot = 0
    
    for record in records:
        if record.attendance_data:
            for day, status in record.attendance_data.items():
                if status == 'P':
                    total_present += 1
                    # Check if it's Sunday for OT calculation
                    if month_filter and year_filter:
                        try:
                            day_date = date(int(year_filter), month_num, int(day))
                            if day_date.weekday() != 6:  # Not Sunday
                                total_normal_ot += 2  # 2 hours normal OT
                        except:
                            pass
                elif status == 'A':
                    total_absent += 1
                elif status == 'H':
                    total_holiday += 1
                elif status == 'M':
                    total_medical += 1
    
    context = {
        'records': records,
        'months': ['January', 'February', 'March', 'April', 'May', 'June',
                  'July', 'August', 'September', 'October', 'November', 'December'],
        'years': range(2020, 2030),
        'days_range': range(1, 32),
        'selected_month': month_filter,
        'selected_year': year_filter,
        'total_present': total_present,
        'total_absent': total_absent,
        'total_holiday': total_holiday,
        'total_medical': total_medical,
        'total_normal_ot': total_normal_ot,
        'total_bonus_ot': total_bonus_ot
    }
    return render(request, 'attendance.html', context)

@custom_staff_required
def delete_attendance_record(request, record_id):
    if request.method == 'POST':
        try:
            record = get_object_or_404(AttendanceRecord, id=record_id)
            record.delete()
            return JsonResponse({'success': True, 'message': 'Record deleted successfully'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def delete_attendance_bulk(request):
    if request.method == 'POST':
        try:
            import json
            data = json.loads(request.body)
            
            # Check if delete all is requested
            if data.get('delete_all'):
                deleted_count = AttendanceRecord.objects.all().delete()[0]
                return JsonResponse({
                    'success': True, 
                    'message': f'Successfully deleted all {deleted_count} attendance records'
                })
            
            # Regular bulk delete
            record_ids = data.get('ids', [])
            
            if not record_ids:
                return JsonResponse({'success': False, 'message': 'No records selected'})
            
            deleted_count = AttendanceRecord.objects.filter(id__in=record_ids).delete()[0]
            return JsonResponse({
                'success': True, 
                'message': f'Successfully deleted {deleted_count} records'
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid request method'})

@custom_staff_required
def delete_all_attendance(request):
    """Delete all attendance records"""
    if request.method == 'POST':
        try:
            deleted_count = AttendanceRecord.objects.all().delete()[0]
            return JsonResponse({
                'success': True, 
                'message': f'Successfully deleted all {deleted_count} attendance records'
            })
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)})
    return JsonResponse({'success': False, 'message': 'Invalid request method'})